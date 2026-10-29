#!/usr/bin/env python3
"""Auto-refresh the dashboard's core 2000-present data (crime + DV) from NYPD's published files.

Downloads NYPD "Historical Crime Data by Precinct" Excel files (CompStat) and the annual
Domestic Violence "Complaint & Radio Run" xlsx, re-parses them, carries the population model and
CCRB series forward from core_base.json, rebuilds the bundle, VALIDATES it, and writes core_data.json
(the dashboard loads it, falling back to its baked snapshot if absent/invalid).

Runs in the GitHub Action (which can reach nyc.gov). pip install: requests xlrd openpyxl numpy
"""
import requests, re, os, sys, json, glob, tempfile, datetime
import numpy as np
HERE=os.path.dirname(os.path.abspath(__file__))
NYC='https://www.nyc.gov'
HIST='https://www.nyc.gov/site/nypd/stats/crime-statistics/historical.page'
PATS={'seven':'seven-major','nonseven':'non-seven-major','misd':'misdemeanor','viol':'violation'}
CLASS={'seven':'Seven Major Felony','nonseven':'Non-Seven Major Felony','misd':'Misdemeanor','viol':'Violation'}
TODAY=datetime.date.today()

def log(*a): print(*a, flush=True)

def crime_urls():
    urls={}
    try:
        html=requests.get(HIST, timeout=120, headers={'User-Agent':'Mozilla/5.0'}).text
        for k,p in PATS.items():
            m=re.search(r'href="([^"]*'+p+r'-offenses-by-precinct-2000-\d+\.xls)"', html, re.I)
            if m: u=m.group(1); urls[k]=u if u.startswith('http') else NYC+u
    except Exception as e:
        log('warn: scrape failed', e)
    # fallback: construct for the most recent plausible year
    base=NYC+'/assets/nypd/downloads/excel/analysis_and_planning/historical-crime-data/'
    for k,p in PATS.items():
        if k in urls: continue
        for y in range(TODAY.year, 2019, -1):
            u=f'{base}{p}-offenses-by-precinct-2000-{y}.xls'
            try:
                if requests.head(u, timeout=30, headers={'User-Agent':'Mozilla/5.0'}).status_code==200: urls[k]=u; break
            except: pass
    return urls

def download(url, path):
    r=requests.get(url, timeout=180, headers={'User-Agent':'Mozilla/5.0'}); r.raise_for_status()
    open(path,'wb').write(r.content); return path

def parse_xls(path):
    import xlrd
    sh=xlrd.open_workbook(path).sheet_by_index(0)
    hr=None; ycols=[]
    for r in range(min(10,sh.nrows)):
        vals=[sh.cell_value(r,c) for c in range(sh.ncols)]
        yrs=[(c,int(float(v))) for c,v in enumerate(vals) if str(v).strip().replace('.0','').isdigit() and 1999<int(float(v))<2100]
        if len(yrs)>=10: hr=r; ycols=yrs; break
    years=[y for _,y in ycols]; cols=[c for c,_ in ycols]; data={}; cur=None
    for r in range(hr+1, sh.nrows):
        a=sh.cell_value(r,0); b=str(sh.cell_value(r,1)).strip()
        if isinstance(a,float) and a>0: cur=int(a)
        elif isinstance(a,str) and a.strip():
            t=a.strip(); cur=int(float(t)) if t.replace('.0','').isdigit() else None
        if cur is None or not b: continue
        bn=' '.join(b.split()); key='TOTAL' if bn.upper().startswith('TOTAL') else bn.upper()
        vals=[]
        for c in cols:
            v=sh.cell_value(r,c)
            try: vals.append(int(float(v)) if v!='' else 0)
            except: vals.append(0)
        data.setdefault(cur,{})[key]=vals
    return years, data

def dv_cols(ws):
    for r in range(1,15):
        row=[re.sub(r'\s+',' ',str(ws.cell(r,c).value or '').strip().lower()) for c in range(1,ws.max_column+1)]
        if 'precinct' in row:
            idx={}
            for c,v in enumerate(row,1):
                if v=='precinct': idx['p']=c
                elif v=='radio runs': idx['radio']=c
                elif v=='rape complaints': idx['rape']=c
                elif v=='felony assault complaints': idx['fa']=c
                elif v=='murder complaints': idx['mur']=c
                elif 'percent total felony' in v: idx['pdom']=c
                elif 'percent domestic violence complaints involving intimate' in v: idx['pip']=c
            return r, idx
    return None, None

def build():
    base=json.load(open(os.path.join(HERE,'core_base.json')))
    RESI=[int(p) for p in base['precincts']]; YEARS=base['years']
    pop={int(p):base['pop'][p] for p in base['pop']}
    def popY(p,i): a=pop[p]; return a[i] if i<len(a) else a[-1]
    tmp=tempfile.mkdtemp()
    urls=crime_urls(); log('crime urls:', {k:urls.get(k) for k in PATS})
    parsed={}
    for k in PATS:
        if k not in urls: raise RuntimeError(f'missing crime url for {k}')
        parsed[k]=parse_xls(download(urls[k], os.path.join(tmp,k+'.xls')))
    years=parsed['seven'][0]
    if not all(parsed[k][0]==years for k in PATS): raise RuntimeError('year columns differ across crime files')
    YEARS=years; NY=len(YEARS); allp=sorted(parsed['seven'][1].keys())
    log('parsed years', YEARS[0],'-',YEARS[-1],'precincts', len(allp))
    counts={}
    def put(m,p,v): counts.setdefault(m,{})[p]=v
    for pct in allp:
        sv=parsed['seven'][1].get(pct,{})
        for off,vals in sv.items():
            if off=='TOTAL': put('Seven Major Felony',pct,vals)
            else: put(off.title(),pct,vals)
        def g(n): return sv.get(n,[0]*NY)
        put('Violent (index crimes)',pct,[a+b+c+d for a,b,c,d in zip(g('MURDER & NON NEGL. MANSLAUGHTER'),g('RAPE'),g('ROBBERY'),g('FELONY ASSAULT'))])
        put('Property (index crimes)',pct,[a+b+c for a,b,c in zip(g('BURGLARY'),g('GRAND LARCENY'),g('GRAND LARCENY OF MOTOR VEHICLE'))])
        tot=[0]*NY
        for k in PATS:
            d=parsed[k][1].get(pct,{})
            for off,vals in d.items():
                if off=='TOTAL': put(CLASS[k],pct,vals); tot=[a+b for a,b in zip(tot,vals)]
                elif k!='seven': put(off.title(),pct,vals)
        put('All Reported Crime',pct,tot)
    # metric order
    ORDER=[('All Reported Crime','headline'),('Violent (index crimes)','group'),('Property (index crimes)','group'),
           ('Seven Major Felony','class'),('Non-Seven Major Felony','class'),('Misdemeanor','class'),('Violation','class')]
    for m in ['Murder & Non Negl. Manslaughter','Rape','Robbery','Felony Assault','Burglary','Grand Larceny','Grand Larceny Of Motor Vehicle']: ORDER.append((m,'off7'))
    placed=set(m for m,_ in ORDER)
    def offs(k,g):
        out=[]
        for pct,o in parsed[k][1].items():
            for raw in o:
                if raw=='TOTAL': continue
                nm=' '.join(raw.split()).title()
                if nm not in placed and nm not in out: out.append(nm)
        return [(m,g) for m in out]
    ORDER+=offs('nonseven','offNS')+offs('misd','offMI')+offs('viol','offVI')
    def series(m,p): return counts.get(m,{}).get(p,[0]*NY)
    cityCrime={}
    for m,_ in ORDER:
        city=[];q1=[];md=[];q3=[]
        for i in range(NY):
            rates=np.array([series(m,p)[i]/popY(p,i)*1e5 for p in RESI])
            t=sum(series(m,p)[i] for p in RESI); ps=sum(popY(p,i) for p in RESI)
            city.append(round(t/ps*1e5,1)); a,b,c=np.percentile(rates,[25,50,75]); q1.append(round(a,1));md.append(round(b,1));q3.append(round(c,1))
        cityCrime[m]={'city':city,'q1':q1,'med':md,'q3':q3}
    # DV
    import openpyxl
    dv={}; dvYears=[]
    for y in range(2020, TODAY.year+1):
        u=f'{NYC}/assets/nypd/downloads/excel/domestic-violence/{y}/dv-complaint-radio-run-{y}.xlsx'
        try:
            p=download(u, os.path.join(tmp,f'dv{y}.xlsx'))
        except Exception: continue
        wb=openpyxl.load_workbook(p,data_only=True); ws=wb[wb.sheetnames[0]]
        hr,idx=dv_cols(ws)
        if not idx: continue
        dvYears.append(y)
        for r in range(hr+1, ws.max_row+1):
            pr=ws.cell(r,idx['p']).value; mm=re.match(r'^\s*0*(\d{1,3})\s*$', str(pr) if pr is not None else '')
            if not mm: continue
            P=int(mm.group(1))
            def gv(key):
                v=ws.cell(r,idx[key]).value if key in idx else None
                try: return float(v)
                except: return 0.0
            dv.setdefault(str(P),{})[str(y)]=[int(gv('radio')),int(gv('rape')),int(gv('fa')),int(gv('mur')),round(gv('pdom'),4),round(gv('pip'),4)]
    dvYears=sorted(set(dvYears)) or base['dvYears']
    if not dv: dv=base['dv']
    dvCity={'radio':[],'radio_q1':[],'radio_q3':[],'fa':[]}
    for y in dvYears:
        i=YEARS.index(y) if y in YEARS else NY-1
        rr=[dv[str(p)][str(y)][0]/popY(p,i)*1000 for p in RESI if str(p) in dv and str(y) in dv[str(p)]]
        if not rr: rr=[0]
        tr=sum(dv[str(p)][str(y)][0] for p in RESI if str(p) in dv and str(y) in dv[str(p)]); tfa=sum(dv[str(p)][str(y)][2] for p in RESI if str(p) in dv and str(y) in dv[str(p)]); ps=sum(popY(p,i) for p in RESI)
        dvCity['radio'].append(round(tr/ps*1000,2)); dvCity['fa'].append(round(tfa/ps*1e5,1))
        dvCity['radio_q1'].append(round(np.percentile(rr,25),2)); dvCity['radio_q3'].append(round(np.percentile(rr,75),2))
    core={'years':YEARS,'precincts':RESI,'pop':{str(p):pop[p] for p in RESI},
          'metrics':[{'id':m,'group':g} for m,g in ORDER],
          'counts':{m:{str(p):series(m,p) for p in RESI} for m,_ in ORDER},
          'cityCrime':cityCrime,'dvYears':dvYears,'dv':dv,'dvCity':dvCity,
          'ccrb':base['ccrb'],'ccrbCity':base['ccrbCity'],'focus':114}
    # validate
    msg=[]
    if len(core['precincts'])!=len(RESI): msg.append('precincts')
    if len(core['metrics'])!=41: msg.append(f"metrics {len(core['metrics'])}")
    if core['years'][0]!=2000 or core['years'][-1]<2023: msg.append('years')
    ac=core['counts'].get('All Reported Crime',{}).get('114',[])
    if not ac or not (5000<ac[-1]<30000): msg.append(f"114 all-crime {ac[-1] if ac else None}")
    if msg:
        log('VALIDATION FAILED, keeping existing core_data.json:', msg); return None
    return core

if __name__=='__main__':
    try:
        core=build()
    except Exception as e:
        log('refresh_core error (keeping existing core_data.json):', e); sys.exit(0)
    if core:
        json.dump(core, open(os.path.join(HERE,'core_data.json'),'w'), separators=(',',':'))
        log('wrote core_data.json: metrics',len(core['metrics']),'precincts',len(core['precincts']),'years',core['years'][0],'-',core['years'][-1],'dvYears',core['dvYears'])
